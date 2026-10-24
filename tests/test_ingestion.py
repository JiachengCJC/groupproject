"""Real parser fixtures plus Ollama contracts; optional integrations are explicit."""
import gzip
import io
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

from fastapi.testclient import TestClient

from app.config import Settings
from app.database import Database
from app.engine import Engine
from app.ingestion import FileParser, ingest
from app.main import create_app
from app.providers import AgentFindings, DecisionBrief, ImageExtraction, ProviderError, Providers, QueryPlan
from app.retrieval import retrieve


class FileTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.settings = Settings(db_path=Path(self.tmp.name)/"db.sqlite3")
        self.db = Database(self.settings.db_path)
        self.db.initialize()
        self.providers = Providers(self.settings)
        self.parser = FileParser(self.settings, self.providers)

    def tearDown(self):
        self.tmp.cleanup()

    def test_text_html_json_yaml_xml_rtf_and_gzip(self):
        fixtures = {
            "notes.txt": b"Customer retention requires faster onboarding.",
            "notes.unknown": "\u5ba2\u6237\u6d41\u5931\u9700\u8981\u6539\u5584\u670d\u52a1".encode(),
            "page.html": b"<script>secret</script><h1>Retention</h1><p>Improve support.</p>",
            "data.json": b'[{"name":"Core","revenue":12},{"name":"Plus","revenue":20}]',
            "data.jsonl": b'{"name":"Core","revenue":12}\n{"name":"Plus","revenue":20}',
            "policy.yaml": b"policy: improve customer support\n",
            "policy.xml": b"<policy><rule>Improve support</rule></policy>",
            "policy.rtf": b"{\\rtf1\\ansi Improve customer support.}",
            "policy.txt.gz": gzip.compress(b"Improve customer support."),
        }
        for name, data in fixtures.items():
            with self.subTest(format=name):
                parsed = self.parser.parse(name, data)
                self.assertTrue(parsed.segments)
        self.assertNotIn("secret", self.parser.parse("page.html", fixtures["page.html"]).segments[0]["content"])

    def test_pdf_pages_and_original_roundtrip(self):
        import pymupdf
        with pymupdf.open() as doc:
            doc.new_page().insert_text((72, 72), "Customer retention: improve onboarding and customer support.")
            doc.new_page().insert_text((72, 72), "Expansion requires a controlled pilot and validation of regional demand.")
            content = doc.tobytes()
        result = ingest(self.db, self.settings, self.providers, "research.pdf", content)
        self.assertEqual(self.db.upload(result["file_id"])["content"], content)
        hits = retrieve(self.db, "regional demand", document_ids=[result["documents"][0]["id"]])
        self.assertEqual(hits[0]["location"], "Page 2")
        self.assertEqual(hits[0]["file_id"], result["file_id"])
        duplicate = ingest(self.db, self.settings, self.providers, "research.pdf", content)
        self.assertEqual(duplicate["status"], "duplicate")

    def test_docx_paragraphs_and_tables(self):
        from docx import Document
        doc = Document()
        doc.add_heading("Customer interviews", 1)
        doc.add_paragraph("Customers report slow onboarding.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Issue"
        table.cell(0, 1).text = "Count"
        table.cell(1, 0).text = "Support delays"
        table.cell(1, 1).text = "12"
        out = io.BytesIO();doc.save(out)
        parsed = self.parser.parse("interviews.docx", out.getvalue())
        self.assertIn("Support delays", "\n".join(s["content"] for s in parsed.segments))
        self.assertTrue(any("Table" in s["location"] for s in parsed.segments))

    def test_pptx_and_xlsx_sheets(self):
        from pptx import Presentation
        from openpyxl import Workbook
        deck = Presentation();slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Market opportunities"
        slide.placeholders[1].text = "Validate demand with a controlled pilot."
        out = io.BytesIO();deck.save(out)
        parsed = self.parser.parse("market.pptx", out.getvalue())
        self.assertEqual(parsed.segments[0]["location"], "Slide 1")
        wb = Workbook();wb.active.title = "Sales"
        wb.active.append(["Product", "Revenue"]);wb.active.append(["Core", 12])
        sheet = wb.create_sheet("Customers");sheet.append(["Segment", "Count"]);sheet.append(["Consumer", 30])
        out = io.BytesIO();wb.save(out)
        result = ingest(self.db, self.settings, self.providers, "business.xlsx", out.getvalue())
        self.assertEqual(len(result["datasets"]), 2)
        self.assertEqual(result["datasets"][0]["row_count"], 1)

    def test_cjk_headers_and_retrieval(self):
        result = ingest(self.db, self.settings, self.providers, "\u5ba2\u6237.csv", "\u5ba2\u6237,\u6d41\u5931\u4eba\u6570\n\u4f01\u4e1a\u5ba2\u6237,12\n".encode())
        columns = result["datasets"][0]["columns"]
        self.assertEqual(columns[0]["name"], "col_1")
        self.assertEqual(columns[0]["label"], "\u5ba2\u6237")
        d = self.db.add_document("\u8bbf\u8c08", "\u5ba2\u6237\u6d41\u5931\u589e\u52a0，\u5ba2\u6237\u53cd\u6620\u670d\u52a1\u54cd\u5e94\u65f6\u95f4\u8fc7\u957f。", "\u8bbf\u8c08")
        self.assertTrue(retrieve(self.db, "\u4e3a\u4ec0\u4e48\u5ba2\u6237\u6d41\u5931\u589e\u52a0？", document_ids=[d["id"]]))

    def test_zip_partial_results_and_no_path_extraction(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            z.writestr("../../policy.txt", "Customer retention policy.")
            z.writestr("bad.exe", b"\x00\x01\x02")
        parsed = self.parser.parse("bundle.zip", out.getvalue())
        self.assertIn("policy.txt", parsed.segments[0]["location"])
        self.assertTrue(parsed.warnings)

    def test_epub_spine_and_odt(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            z.writestr("META-INF/container.xml", '<container><rootfile full-path="book/content.opf"/></container>')
            z.writestr("book/content.opf", '<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="c" href="chapter.html"/></manifest><spine><itemref idref="c"/></spine></package>')
            z.writestr("book/chapter.html", "<h1>Customer retention</h1><p>Improve support.</p>")
        self.assertEqual(self.parser.parse("book.epub", out.getvalue()).segments[0]["location"], "Chapter 1")
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            z.writestr("content.xml", '<root xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><text:p>Customer retention policy</text:p></root>')
        self.assertTrue(self.parser.parse("policy.odt", out.getvalue()).segments)

    def test_email_body_and_attachment(self):
        from email.message import EmailMessage
        msg = EmailMessage();msg["Subject"] = "Customer research";msg.set_content("Customers need faster support.")
        msg.add_attachment(b"Improve onboarding.", maintype="text", subtype="plain", filename="notes.txt")
        parsed = self.parser.parse("research.eml", msg.as_bytes())
        self.assertTrue(any("Attachment" in s["location"] for s in parsed.segments))

    def test_parquet_to_queryable_table(self):
        import duckdb
        file = Path(self.tmp.name)/"input.parquet"
        with duckdb.connect() as con:
            con.execute("COPY (SELECT 'Core' AS product, 12 AS revenue) TO ? (FORMAT PARQUET)", [str(file)])
        self.assertEqual(self.parser.parse(file.name, file.read_bytes()).tables[0]["name"], "input")

    def test_reject_binary_encryption_and_expansion(self):
        for name, data in [("bad.exe", b"MZ\x00\x01"), ("empty.txt", b"")]:
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.parser.parse(name, data)
        import pymupdf
        with pymupdf.open() as doc:
            doc.new_page().insert_text((72,72), "Secret policy document.")
            data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
        with self.assertRaisesRegex(ValueError, "encrypted"):
            self.parser.parse("locked.pdf", data)

    @unittest.skipUnless(shutil.which("tesseract"), "Tesseract not installed")
    def test_image_ocr_and_scanned_pdf(self):
        from PIL import Image, ImageDraw, ImageFont
        import pymupdf
        image = Image.new("RGB", (1400, 200), "white")
        fontpath = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        if not Path(fontpath).exists():
            fontpath = "/System/Library/Fonts/Supplemental/Arial.ttf"
        ImageDraw.Draw(image).text((30,50), "Customer retention policy", fill="black", font=ImageFont.truetype(fontpath, 48))
        out = io.BytesIO();image.save(out, format="PNG")
        self.assertIn("retention", self.parser.parse("scan.png", out.getvalue()).segments[0]["content"].lower())
        with pymupdf.open() as doc:
            page = doc.new_page(width=700, height=100);page.insert_image(page.rect, stream=out.getvalue())
            data = doc.tobytes()
        self.assertIn("Page 1", self.parser.parse("scan.pdf", data).segments[0]["location"])

    def test_vision_is_opt_in_and_marked(self):
        from PIL import Image
        out = io.BytesIO();Image.new("RGB", (100,100), "white").save(out, format="PNG")
        parser = FileParser(self.settings, self.providers, vision_provider="ollama", vision_model="vision-test")
        with patch.object(self.providers, "structured", return_value=({"text":"Revenue 12", "description":"A chart", "limitations":["Check numbers"]}, {})) as call:
            result = parser.parse("chart.png", out.getvalue())
        self.assertEqual(call.call_args.kwargs["provider"], "ollama")
        self.assertEqual(result.method, "ollama-vision")

    def test_media_has_actionable_dependency_message(self):
        with self.assertRaisesRegex(ValueError, "WHISPER_MODEL"):
            self.parser.parse("meeting.mp3", b"audio")

    def test_invalid_archive_and_oversized_table_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "parse"):
            self.parser.parse("broken.zip", b"PK not a zip file")
        with self.assertRaisesRegex(ValueError, "10,000"):
            self.parser.parse("large.csv", b"name,value\n"+b"x,1\n"*10001)
        with self.assertRaises(ValueError):
            self.parser.parse("entities.xml", b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>')

    def test_pdf_blank_page_does_not_discard_readable_pages(self):
        import pymupdf
        with pymupdf.open() as doc:
            doc.new_page().insert_text((72,72), "Customer research: improve support and customer retention.")
            doc.new_page()
            data=doc.tobytes()
        with patch.object(self.parser,"image",side_effect=ValueError("No text detected")):
            result=self.parser.parse("mixed.pdf",data)
        self.assertEqual(len(result.segments),1)
        self.assertIn("Page 2",result.warnings[0])

    @unittest.skipUnless(shutil.which("libreoffice") or shutil.which("soffice"), "LibreOffice not installed")
    def test_legacy_word_and_powerpoint_conversion(self):
        from docx import Document
        from pptx import Presentation
        folder=Path(self.tmp.name)
        word=Document();word.add_paragraph("Customer retention policy: improve support.");word.save(folder/"fixture.docx")
        slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text="Customer retention";slides.save(folder/"fixture.pptx")
        binary=shutil.which("libreoffice") or shutil.which("soffice")
        for source,target in [("docx","doc"),("pptx","ppt")]:
            process=subprocess.run([binary,"-env:UserInstallation="+(folder/"profile").as_uri(),"--headless","--convert-to",target,"--outdir",str(folder),str(folder/("fixture."+source))],capture_output=True,timeout=90)
            self.assertEqual(process.returncode,0)
            file=folder/("fixture."+target)
            result=self.parser.parse(file.name,file.read_bytes())
            self.assertIn("Customer retention", "\n".join(s["content"] for s in result.segments))

    def test_optional_whisper_timestamp_contract(self):
        import sys
        from types import SimpleNamespace
        settings=Settings(db_path=self.settings.db_path,whisper_model="local-fixture")
        model=Mock();model.transcribe.return_value=([SimpleNamespace(start=0,end=2.5,text="Customer support is slow.")],SimpleNamespace(duration=3))
        module=SimpleNamespace(WhisperModel=Mock(return_value=model))
        with patch.dict(sys.modules,{"faster_whisper":module}):
            parsed=FileParser(settings,self.providers).parse("meeting.wav",b"fixture audio")
        self.assertEqual(parsed.segments[0]["location"],"0.0s–2.5s")
        self.assertEqual(parsed.method,"whisper")

    def test_document_only_local_and_ollama_pipeline(self):
        d = self.db.add_document("Interviews", "Customers report slow support and onboarding. Improve customer retention.", "test")
        request = {"question":"How should we improve customer retention?", "document_ids":[d["id"]], "agent_mode":"local"}
        engine = Engine(self.db, self.settings)
        local = engine.analyze(request)
        self.assertEqual(local["metrics"], {})
        self.assertTrue(all(e["kind"] == "document" for e in local["evidence"]))
        def structured(schema, instructions, context, **kwargs):
            self.assertEqual(kwargs["provider"], "ollama")
            ids = [e["id"] for e in context["evidence"]]
            if schema is AgentFindings:
                return {"summary":"Customers report delays.", "claims":[{"text":"Customers report delays.", "evidence_ids":ids[:1]}], "limitations":[]}, {}
            self.assertIs(schema, DecisionBrief)
            return {"recommendation":"Pilot faster onboarding", "explanation":"The interviews report delays.", "alternatives":["Measure response times"], "evidence_ids":ids[:1]}, {}
        with patch.object(engine.providers, "structured", side_effect=structured):
            result = engine.analyze({**request, "agent_mode":"ollama", "model":"test-model", "decision_mode":"llm"})
        self.assertEqual(result["decision"]["provider"], "ollama")
        self.assertIsNone(result["decision"]["scores"])
        self.assertEqual(result["measurement"]["citation_coverage"], 1)


class OllamaTests(unittest.TestCase):
    def test_json_schema_and_token_usage(self):
        response = Mock()
        response.json.return_value = {"done":True, "message":{"content":json.dumps({"sql":"SELECT 1", "search_query":"retention"})}, "prompt_eval_count":12, "eval_count":8}
        providers = Providers(Settings(ollama_model="local-test"))
        with patch("app.providers.httpx.post", return_value=response) as post:
            data, usage = providers.structured(QueryPlan, "Read-only SQL", {}, provider="ollama")
        self.assertEqual(data["sql"], "SELECT 1")
        self.assertEqual(post.call_args.kwargs["json"]["format"], QueryPlan.model_json_schema())
        self.assertFalse(post.call_args.kwargs["json"]["stream"])
        self.assertEqual(usage["output_tokens"], 8)

    def test_malformed_or_incomplete_response_fails(self):
        providers = Providers(Settings(ollama_model="test"))
        for payload in [{"done":True,"message":{"content":"{}"}}, {"done":False,"message":{"content":"{}"}}]:
            response = Mock();response.json.return_value = payload
            with patch("app.providers.httpx.post", return_value=response), self.assertRaises(ProviderError):
                providers.structured(QueryPlan, "Test", {}, provider="ollama")


class UploadAPITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.client = TestClient(create_app(Settings(db_path=Path(self.tmp.name)/"api.db")))
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None,None,None);self.tmp.cleanup()

    def test_upload_preview_original_and_duplicate(self):
        files={"file":("policy.txt", b"Customer retention policy", "text/plain")}
        response = self.client.post("/api/uploads", files=files, data={"category":"Internal policy"})
        self.assertEqual(response.status_code, 201, response.text)
        result = response.json()
        self.assertEqual(self.client.get("/api/files/"+result["file_id"]).content, b"Customer retention policy")
        doc = self.client.get("/api/documents/"+result["documents"][0]["id"]).json()
        self.assertEqual(doc["category"], "Internal policy")
        self.assertEqual(self.client.post("/api/uploads",files=files).json()["status"], "duplicate")

    def test_scope_validation_and_multipart_origin(self):
        self.assertEqual(self.client.post("/api/analyses", json={"question":"Analyze these documents"}).status_code, 400)
        d = self.client.post("/api/documents", json={"title":"Private scope","content":"Customer policy","dataset_id":"demo_retention"}).json()
        self.assertEqual(self.client.post("/api/analyses", json={"question":"Analyze these documents", "dataset_id":"demo_pricing", "document_ids":[d["id"]]}).status_code, 400)
        files={"file":("a.txt", b"policy")}
        self.assertEqual(self.client.post("/api/uploads",files=files, headers={"origin":"https://example.org"}).status_code, 403)
        self.assertEqual(self.client.post("/api/uploads",files={"file":("bad.exe", b"MZ\x00")}).status_code, 400)
        self.assertEqual(self.client.post("/api/uploads",files=files,data={"vision_provider":"ollama"}).status_code, 400)

    def test_ollama_model_validation(self):
        with patch.object(self.client.app.state.engine.providers,"ollama_models",return_value={"available":True,"models":["local-test"]}):
            response=self.client.post("/api/analyses",json={"question":"Analyze our performance", "dataset_id":"demo_pricing", "agent_mode":"ollama","model":"missing"})
            self.assertEqual(response.status_code,400)


if __name__ == "__main__":
    unittest.main()
